import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx2pdf import convert
import pypdf

def update_report():
    doc = docx.Document("report/Mess report_backup.docx")

    # 1. Update content paragraphs font size to strictly 12pt, Times New Roman, Justified, 1.5 line spacing
    def style_content_paragraph(p, text=None):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(12)
        if text is not None:
            p.text = text
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 0, 0)

    # Abstract paragraphs [27] to [30]
    abstract_texts = [
        "The Hostel Mess Management System is a web-based application designed to streamline and automate hostel mess dining operations by providing an integrated platform where students and mess administrators can efficiently interact.",
        "Through the portal, students can view weekly cyclical dining menus, daily standard rates, opt into scheduled additional food items at least one day in advance, manage cancellation requests with valid reasons prior to 12 hours of meal times, track real-time monthly mess dues, and submit queries or grievances seamlessly.",
        "The system's functionality is driven by five key modules: User & Role Management for secure access (Administrator and Student accounts), Weekly Cyclical Menu Scheduling with standard daily fee control, Additional Food Scheduling with advance opt-in and reason-based cancellation workflows, Automated Monthly Fee Calculation integrating base dining dues and additional surcharges, and a Student Query Resolution Portal. Built on a lightweight modern tech stack featuring Python Flask, MongoDB, and Vanilla CSS, the application emphasizes high performance, zero external styling bloat, and transparent hostel dining administration.",
        "By eliminating manual bottlenecks, paper coupons, and billing disputes, the Hostel Mess Management System empowers hostel administration with a centralized, automated, and efficient platform that significantly enhances student dining satisfaction and simplifies mess operations."
    ]
    for idx, txt in zip(range(27, 31), abstract_texts):
        style_content_paragraph(doc.paragraphs[idx], txt)

    # 1.1 Introduction paragraphs [35] to [37]
    intro_texts = [
        "Hostel dining operations play a vital role in shaping the everyday living experience of residential college students, providing them with nutritious and timely meals throughout their academic tenure. Traditionally, the management of hostel mess services has been conducted through manual logbooks, paper meal coupons, and disjointed verbal communication between students, catering staff, and hostel wardens.",
        "To address these operational bottlenecks, the Hostel Mess Management System has been conceptualized and implemented as a centralized, automated web application. The platform is designed to digitize the complete hostel dining lifecycle, starting from role-based authentication, weekly meal menu publishing, and daily fee tracking to advance scheduling of additional food items, automated billing, and transparent grievance resolution.",
        "The system incorporates automated rule validation including advance day opt-in requirements, reason-based cancellation checks, and dynamic fee calculation, ensuring an equitable, transparent, and hassle-free dining environment for residential students and administrators alike."
    ]
    for idx, txt in zip(range(35, 38), intro_texts):
        style_content_paragraph(doc.paragraphs[idx], txt)

    # 1.2 Problem Statement paragraphs [40] to [41]
    problem_texts = [
        "The existing hostel dining administration across many colleges is largely manual, paper-intensive, and fragmented, resulting in inefficient communication among students, mess managers, and billing accountants. This creates significant difficulty in planning daily kitchen procurement and meal preparation quantities, as administrators have no accurate means of forecasting student attendance or demand for special dishes in advance.",
        "Furthermore, the reliance on handwritten paper tokens and manual register entries is highly susceptible to human error, loss of records, and billing discrepancies at the end of each billing cycle. Students often experience a lack of transparency regarding additional food charges, and last-minute cancellations without justification cause substantial food waste and financial loss for the institution. Concurrently, students lack a structured, trackable grievance portal to register feedback regarding hygiene, menu quality, or service. Therefore, there is a clear and urgent need for an automated system to manage the entire hostel dining process in a structured, efficient, and transparent manner."
    ]
    for idx, txt in zip(range(40, 42), problem_texts):
        style_content_paragraph(doc.paragraphs[idx], txt)

    # 1.3 Objectives paragraphs [44] to [46]
    objective_texts = [
        "The primary objective of the Hostel Mess Management System is to establish a centralized digital platform that efficiently manages all mess-related operations. It provides an intuitive interface for students to view weekly dining schedules, daily standard fees, and opt into additional food items, while simultaneously enabling mess administrators to manage menus, schedule specials, and process student queries.",
        "The system enforces automated business rules including requiring additional dishes to be accepted at least one day in advance, disallowing same-day orders, and implementing a 12-hour cancellation policy where cancellations before 12 hours require valid reasons and administrator approval, while cancellations within 12 hours are strictly prohibited to prevent kitchen waste.",
        "To ensure financial accuracy and transparency, the platform computes dynamic monthly mess fees incorporating both standard daily rates and accepted additional surcharges. Additionally, it provides an integrated queries and grievance resolution portal to foster continuous quality improvement and student satisfaction."
    ]
    for idx, txt in zip(range(44, 47), objective_texts):
        style_content_paragraph(doc.paragraphs[idx], txt)

    # Conclusion paragraphs [437] to [441]
    conclusion_texts = [
        "The Hostel Mess Management System has been successfully designed and developed to automate and simplify the traditional, manual hostel dining administration. The system integrates students and mess administrators onto a single digital platform, ensuring smooth interaction and efficient management of all dining-related activities. By leveraging modern technologies like Python Flask for a lightweight, secure backend, MongoDB for dynamic document persistence, and Vanilla CSS for responsive, zero-dependency styling, the portal delivers a reliable and modern user experience.",
        "The implementation of automated business logic such as advance-day opt-in enforcement, 12-hour cancellation rules with mandatory reason verification, and dynamic monthly fee aggregation has dramatically enhanced the efficiency, transparency, and accuracy of mess operations. Students can effortlessly inspect cyclical menus, opt into desired extra dishes, track real-time billing, and submit grievances with immediate status visibility. Concurrently, administrators benefit from centralized menu control, automated headcount tracking for special meals, and a streamlined workflow for reviewing student rejection requests.",
        "Through rigorous practical implementation and testing, the portal demonstrated substantial improvements in operational accuracy and communication flow. It successfully eliminated manual paperwork, paper coupon loss, and billing disputes while curbing unnecessary food wastage by preventing unverified last-minute cancellations.",
        "In conclusion, the Hostel Mess Management System provides a comprehensive, scalable, and practical solution that transforms conventional mess management into an automated, accountable platform. It ensures fairness, transparency, and operational efficiency across the residential dining lifecycle, establishing a foundation for smarter campus dining administration.",
        "Overall, this project proves that the adoption of modern web frameworks and automated operational workflows can revolutionize residential campus dining management and pave the way for a smarter and more connected institutional ecosystem."
    ]
    for idx, txt in zip(range(437, 442), conclusion_texts):
        style_content_paragraph(doc.paragraphs[idx], txt)

    # 2. Update 2.2 PROJECT CODE section
    # Find paragraph 53 ("2.2 PROJECT CODE") and paragraph 408 (page break before "CHAPTER 3")
    start_p_idx = 54  # "backend/app.py:"
    end_p_idx = 408    # the empty paragraph with page break before CHAPTER 3

    # Gather updated code snippets
    with open("app.py", "r", encoding="utf-8") as f:
        app_code = f.read()

    with open("templates/auth.html", "r", encoding="utf-8") as f:
        auth_code = f.read()

    with open("templates/student_specials.html", "r", encoding="utf-8") as f:
        student_specials_code = f.read()

    with open("templates/admin_specials.html", "r", encoding="utf-8") as f:
        admin_specials_code = f.read()

    with open("templates/student_fees.html", "r", encoding="utf-8") as f:
        student_fees_code = f.read()

    with open("templates/admin_menu.html", "r", encoding="utf-8") as f:
        admin_menu_code = f.read()

    code_sections = [
        ("backend/app.py:", app_code),
        ("templates/auth.html:", auth_code),
        ("templates/student_specials.html:", student_specials_code),
        ("templates/admin_specials.html:", admin_specials_code),
        ("templates/student_fees.html:", student_fees_code),
        ("templates/admin_menu.html:", admin_menu_code)
    ]

    # Build the list of paragraphs to insert
    new_code_items = []
    for heading, src in code_sections:
        new_code_items.append((True, heading))
        for line in src.splitlines():
            clean_line = line.replace("\xa0", " ")
            new_code_items.append((False, clean_line))

    # The original document had (408 - 54) = 354 paragraphs for code
    # We will modify the existing paragraphs in place, and add or remove as necessary
    existing_code_p = doc.paragraphs[start_p_idx:end_p_idx]
    
    # Helper to format a code paragraph
    def format_code_p(p, is_heading, text):
        p.text = text if text else " "
        p.paragraph_format.space_before = Pt(6) if is_heading else Pt(0)
        p.paragraph_format.space_after = Pt(2) if is_heading else Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            if is_heading:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.bold = True
            else:
                r.font.name = "Courier New"
                r.font.size = Pt(8.5)
                r.font.bold = False

    # Reuse existing paragraphs
    min_len = min(len(existing_code_p), len(new_code_items))
    for i in range(min_len):
        is_hdr, text = new_code_items[i]
        format_code_p(existing_code_p[i], is_hdr, text)

    # If new_code_items has more than existing, insert them before P408
    target_p = doc.paragraphs[end_p_idx]
    if len(new_code_items) > len(existing_code_p):
        for i in range(min_len, len(new_code_items)):
            is_hdr, text = new_code_items[i]
            new_p = target_p.insert_paragraph_before()
            format_code_p(new_p, is_hdr, text)
    elif len(existing_code_p) > len(new_code_items):
        # Clear unused existing paragraphs
        for i in range(min_len, len(existing_code_p)):
            p = existing_code_p[i]
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0

    # Save to report/Mess report.docx
    out_docx = os.path.abspath("report/Mess report.docx")
    out_pdf = os.path.abspath("report/Mess report.pdf")
    doc.save(out_docx)
    print("Saved:", out_docx)

    # Convert to PDF
    try:
        convert(out_docx, out_pdf)
        print("Converted:", out_pdf)
        reader = pypdf.PdfReader(out_pdf)
        print("PDF Page count:", len(reader.pages))
    except Exception as e:
        print("PDF conversion note:", e)

if __name__ == "__main__":
    update_report()
